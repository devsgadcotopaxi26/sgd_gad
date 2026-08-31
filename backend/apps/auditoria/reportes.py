"""
Generación de reportes PDF (WeasyPrint) y Excel (openpyxl)
"""
import io
from datetime import datetime
from django.http import HttpResponse
from django.utils import timezone
from weasyprint import HTML, CSS


def generar_pdf(html_content: str, nombre_archivo: str) -> HttpResponse:
    pdf_file = io.BytesIO()
    HTML(string=html_content).write_pdf(pdf_file)
    pdf_file.seek(0)
    response = HttpResponse(pdf_file.read(), content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{nombre_archivo}"'
    return response


def html_base(titulo: str, subtitulo: str, contenido: str) -> str:
    fecha = timezone.now().strftime('%d/%m/%Y %H:%M')
    return f"""
<!DOCTYPE html>
<html lang="es">
<head>
<meta charset="UTF-8">
<style>
  @page {{
    size: A4;
    margin: 2cm 1.5cm;
    @bottom-center {{
      content: "Página " counter(page) " de " counter(pages);
      font-size: 9pt;
      color: #888;
    }}
  }}
  * {{ box-sizing: border-box; margin: 0; padding: 0; }}
  body {{ font-family: Arial, sans-serif; font-size: 10pt; color: #222; line-height: 1.5; }}

  .header {{ display: flex; align-items: center; justify-content: space-between; border-bottom: 3px solid #002f6c; padding-bottom: 12px; margin-bottom: 20px; }}
  .header-left {{ }}
  .inst {{ font-size: 14pt; font-weight: bold; color: #002f6c; }}
  .sub-inst {{ font-size: 9pt; color: #666; margin-top: 2px; }}
  .header-right {{ text-align: right; font-size: 8pt; color: #888; }}

  .report-title {{ background: #002f6c; color: #fff; padding: 10px 14px; border-radius: 6px; margin-bottom: 16px; }}
  .report-title h1 {{ font-size: 13pt; font-weight: bold; }}
  .report-title p {{ font-size: 9pt; opacity: .8; margin-top: 3px; }}

  .kpi-row {{ display: flex; gap: 10px; margin-bottom: 16px; }}
  .kpi-box {{ flex: 1; border: 1px solid #e5e7eb; border-radius: 6px; padding: 10px; text-align: center; }}
  .kpi-box .val {{ font-size: 18pt; font-weight: bold; color: #002f6c; }}
  .kpi-box .lbl {{ font-size: 8pt; color: #666; margin-top: 2px; }}

  table {{ width: 100%; border-collapse: collapse; margin-bottom: 16px; font-size: 9pt; }}
  thead tr {{ background: #002f6c; color: #fff; }}
  thead th {{ padding: 7px 8px; text-align: left; font-weight: bold; font-size: 8.5pt; }}
  tbody tr:nth-child(even) {{ background: #f8faff; }}
  tbody td {{ padding: 6px 8px; border-bottom: 1px solid #f0f0f0; }}

  .section-title {{ font-size: 11pt; font-weight: bold; color: #002f6c; border-bottom: 1px solid #e5e7eb; padding-bottom: 4px; margin-bottom: 10px; margin-top: 16px; }}

  .badge {{ display: inline-block; padding: 2px 8px; border-radius: 10px; font-size: 8pt; font-weight: bold; }}
  .badge-blue  {{ background: #e8f1fd; color: #002f6c; }}
  .badge-green {{ background: #f0fdf4; color: #15803d; }}
  .badge-red   {{ background: #fef2f2; color: #dc2626; }}
  .badge-gray  {{ background: #f3f4f6; color: #374151; }}

  .footer-note {{ margin-top: 20px; font-size: 8pt; color: #888; border-top: 1px solid #e5e7eb; padding-top: 8px; }}
</style>
</head>
<body>
  <div class="header">
    <div class="header-left">
      <div class="inst">GAD Provincial de Cotopaxi</div>
      <div class="sub-inst">Gobierno Autónomo Descentralizado de la Provincia de Cotopaxi</div>
    </div>
    <div class="header-right">
      Sistema de Gestión Documental<br>
      Generado: {fecha}
    </div>
  </div>

  <div class="report-title">
    <h1>{titulo}</h1>
    <p>{subtitulo}</p>
  </div>

  {contenido}

  <div class="footer-note">
    Documento generado automáticamente por el SGD — GAD Provincia de Cotopaxi · {fecha}
  </div>
</body>
</html>
"""


def reporte_tramites(tramites, filtros: dict) -> str:
    from apps.tramites.models import Tramite
    hoy = timezone.now().date()

    total      = len(tramites)
    resueltos  = sum(1 for t in tramites if t.estado == 'resuelto')
    pendientes = sum(1 for t in tramites if t.estado in ('ingresado','asignado','en_proceso'))
    vencidos   = sum(1 for t in tramites if t.fecha_limite < hoy and t.estado not in ('resuelto','archivado','rechazado'))

    kpis = f"""
    <div class="kpi-row">
      <div class="kpi-box"><div class="val">{total}</div><div class="lbl">Total trámites</div></div>
      <div class="kpi-box"><div class="val">{resueltos}</div><div class="lbl">Resueltos</div></div>
      <div class="kpi-box"><div class="val">{pendientes}</div><div class="lbl">Pendientes</div></div>
      <div class="kpi-box"><div class="val">{vencidos}</div><div class="lbl">Vencidos</div></div>
    </div>
    """

    ESTADO_BADGE = {
        'ingresado':  'badge-blue',
        'asignado':   'badge-blue',
        'en_proceso': 'badge-blue',
        'resuelto':   'badge-green',
        'rechazado':  'badge-red',
        'archivado':  'badge-gray',
    }

    filas = ''
    for t in tramites:
        badge = ESTADO_BADGE.get(t.estado, 'badge-gray')
        dias  = (t.fecha_limite - hoy).days if t.estado not in ('resuelto','archivado','rechazado') else '—'
        filas += f"""
        <tr>
          <td><strong>{t.numero_tramite}</strong></td>
          <td>{t.asunto[:60]}</td>
          <td>{t.persona.nombre_completo if t.persona_id else (t.firmante_oficio or '—')}</td>
          <td>{(t.unidad_responsable.siglas or t.unidad_responsable.nombre[:15]) if t.unidad_responsable_id else 'Sin asignar'}</td>
          <td>{t.fecha_ingreso.strftime('%d/%m/%Y')}</td>
          <td>{t.fecha_limite.strftime('%d/%m/%Y')}</td>
          <td><span class="badge {badge}">{t.get_estado_display()}</span></td>
          <td>{dias}</td>
        </tr>
        """

    tabla = f"""
    <div class="section-title">Detalle de trámites</div>
    <table>
      <thead>
        <tr>
          <th>N° Trámite</th><th>Asunto</th><th>Ciudadano</th>
          <th>Unidad</th><th>Ingreso</th><th>Límite</th>
          <th>Estado</th><th>Días rest.</th>
        </tr>
      </thead>
      <tbody>{filas}</tbody>
    </table>
    """

    periodo = filtros.get('periodo', 'Todos los períodos')
    return html_base(
        'Reporte de trámites ciudadanos',
        f'Período: {periodo} · Total: {total} registros',
        kpis + tabla
    )


def reporte_documentos(documentos, filtros: dict) -> str:
    total     = len(documentos)
    firmados  = sum(1 for d in documentos if d.fecha_firma)
    borradores = sum(1 for d in documentos if d.estado == 'borrador')
    enviados  = sum(1 for d in documentos if d.estado == 'enviado')

    kpis = f"""
    <div class="kpi-row">
      <div class="kpi-box"><div class="val">{total}</div><div class="lbl">Total documentos</div></div>
      <div class="kpi-box"><div class="val">{firmados}</div><div class="lbl">Firmados</div></div>
      <div class="kpi-box"><div class="val">{enviados}</div><div class="lbl">Enviados</div></div>
      <div class="kpi-box"><div class="val">{borradores}</div><div class="lbl">Borradores</div></div>
    </div>
    """

    ESTADO_BADGE = {
        'borrador':    'badge-gray',
        'en_revision': 'badge-blue',
        'aprobado':    'badge-green',
        'enviado':     'badge-blue',
        'archivado':   'badge-gray',
        'anulado':     'badge-red',
    }

    filas = ''
    for d in documentos:
        badge = ESTADO_BADGE.get(d.estado, 'badge-gray')
        filas += f"""
        <tr>
          <td><strong>{d.numero_documento or '—'}</strong></td>
          <td>{d.tipo_documento.nombre}</td>
          <td>{d.asunto[:55]}</td>
          <td>{d.unidad_origen.siglas or d.unidad_origen.nombre[:15]}</td>
          <td>{d.creado_en.strftime('%d/%m/%Y')}</td>
          <td><span class="badge {badge}">{d.get_estado_display()}</span></td>
        </tr>
        """

    tabla = f"""
    <div class="section-title">Detalle de documentos</div>
    <table>
      <thead>
        <tr>
          <th>N° Documento</th><th>Tipo</th><th>Asunto</th>
          <th>Unidad</th><th>Fecha</th><th>Estado</th>
        </tr>
      </thead>
      <tbody>{filas}</tbody>
    </table>
    """

    periodo = filtros.get('periodo', 'Todos los períodos')
    return html_base(
        'Reporte de documentos institucionales',
        f'Período: {periodo} · Total: {total} registros',
        kpis + tabla
    )


def reporte_kpi_unidades(unidades_data: list) -> str:
    filas = ''
    for u in unidades_data:
        pct = round(u['resueltos'] / u['total'] * 100, 1) if u['total'] > 0 else 0
        filas += f"""
        <tr>
          <td><strong>{u['unidad']}</strong></td>
          <td>{u['total']}</td>
          <td>{u['pendientes']}</td>
          <td>{u['resueltos']}</td>
          <td>{pct}%</td>
        </tr>
        """

    tabla = f"""
    <div class="section-title">KPIs por unidad organizativa</div>
    <table>
      <thead>
        <tr>
          <th>Unidad</th><th>Total trámites</th>
          <th>Pendientes</th><th>Resueltos (mes)</th><th>% Cumplimiento</th>
        </tr>
      </thead>
      <tbody>{filas}</tbody>
    </table>
    """

    fecha = timezone.now().strftime('%B %Y')
    return html_base(
        'Reporte KPIs por unidad',
        f'Período: {fecha}',
        tabla
    )


# ─── Excel helpers ────────────────────────────────────────────────────────────

def _wb_style(ws, headers: list[str]):
    """Aplica cabecera azul institucional y autoajusta columnas."""
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    azul    = '002F6C'
    rojo    = 'DA291C'
    fill_h  = PatternFill('solid', fgColor=azul)
    fill_z  = PatternFill('solid', fgColor='EEF2FF')
    borde   = Border(
        left=Side(style='thin', color='D0D0D0'),
        right=Side(style='thin', color='D0D0D0'),
        top=Side(style='thin', color='D0D0D0'),
        bottom=Side(style='thin', color='D0D0D0'),
    )
    center  = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left    = Alignment(horizontal='left',   vertical='center', wrap_text=True)

    for col_i, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_i, value=h)
        cell.font      = Font(bold=True, color='FFFFFF', size=10)
        cell.fill      = fill_h
        cell.alignment = center
        cell.border    = borde

    ws.row_dimensions[1].height = 20

    for row in ws.iter_rows(min_row=2):
        is_zebra = row[0].row % 2 == 0
        for cell in row:
            cell.border    = borde
            cell.alignment = left
            if is_zebra:
                cell.fill = fill_z

    for col_i in range(1, len(headers) + 1):
        max_len = max(
            (len(str(ws.cell(r, col_i).value or '')) for r in range(1, ws.max_row + 1)),
            default=10
        )
        ws.column_dimensions[get_column_letter(col_i)].width = min(max_len + 4, 40)

    ws.freeze_panes = 'A2'


def _excel_response(wb, nombre: str) -> HttpResponse:
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    resp = HttpResponse(buf.read(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp['Content-Disposition'] = f'attachment; filename="{nombre}"'
    return resp


def excel_tramites(tramites, filtros: dict) -> HttpResponse:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = 'Trámites'

    headers = ['N°', 'Número', 'Tipo', 'Ciudadano', 'Cédula', 'Estado',
               'Unidad responsable', 'Fecha ingreso', 'Fecha límite', 'Días restantes']
    ws.append(headers)

    hoy = timezone.now().date()
    for i, t in enumerate(tramites, 1):
        dias = (t.fecha_limite.date() - hoy).days if t.fecha_limite else ''
        ws.append([
            i,
            t.numero_tramite or '',
            t.tipo_tramite.nombre if t.tipo_tramite else '',
            f'{t.persona.nombres} {t.persona.apellidos}' if t.persona else '',
            t.persona.cedula if t.persona else '',
            t.estado,
            t.unidad_responsable.siglas or t.unidad_responsable.nombre if t.unidad_responsable else '',
            t.fecha_ingreso.strftime('%d/%m/%Y') if t.fecha_ingreso else '',
            t.fecha_limite.strftime('%d/%m/%Y') if t.fecha_limite else '',
            dias,
        ])

    _wb_style(ws, headers)

    # Hoja resumen
    ws2 = wb.create_sheet('Resumen')
    from collections import Counter
    conteo = Counter(t.estado for t in tramites)
    ws2.append(['Estado', 'Cantidad'])
    for k, v in conteo.items():
        ws2.append([k, v])
    ws2.append(['TOTAL', len(tramites)])
    _wb_style(ws2, ['Estado', 'Cantidad'])

    return _excel_response(wb, f'tramites_{timezone.now().strftime("%Y%m%d")}.xlsx')


def excel_documentos(documentos, filtros: dict) -> HttpResponse:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = 'Documentos'

    headers = ['N°', 'Número oficial', 'Tipo', 'Asunto', 'Unidad origen',
               'Estado', 'Firmado', 'Fecha creación', 'Fecha envío']
    ws.append(headers)

    for i, d in enumerate(documentos, 1):
        ws.append([
            i,
            d.numero_documento or '',
            d.tipo_documento.nombre if d.tipo_documento else '',
            (d.asunto or '')[:80],
            d.unidad_origen.siglas or d.unidad_origen.nombre if d.unidad_origen else '',
            d.estado,
            'Sí' if d.firma_bce_info else 'No',
            d.creado_en.strftime('%d/%m/%Y') if d.creado_en else '',
            d.fecha_envio.strftime('%d/%m/%Y') if getattr(d, 'fecha_envio', None) else '',
        ])

    _wb_style(ws, headers)

    ws2 = wb.create_sheet('Por tipo')
    from collections import Counter
    conteo = Counter(d.tipo_documento.nombre if d.tipo_documento else 'Sin tipo' for d in documentos)
    ws2.append(['Tipo de documento', 'Cantidad'])
    for k, v in sorted(conteo.items(), key=lambda x: -x[1]):
        ws2.append([k, v])
    _wb_style(ws2, ['Tipo de documento', 'Cantidad'])

    return _excel_response(wb, f'documentos_{timezone.now().strftime("%Y%m%d")}.xlsx')


def excel_kpi(unidades_data: list, periodo: str) -> HttpResponse:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.chart import BarChart, Reference

    wb = Workbook()
    ws = wb.active
    ws.title = 'KPI por unidad'

    headers = ['Unidad', 'Total trámites', 'Pendientes', 'Resueltos (30d)', '% Cumplimiento']
    ws.append(headers)

    for d in unidades_data:
        pct = round(d['resueltos'] / d['total'] * 100, 1) if d['total'] > 0 else 0
        ws.append([d['unidad'], d['total'], d['pendientes'], d['resueltos'], pct])

    _wb_style(ws, headers)

    # Gráfico de barras
    chart = BarChart()
    chart.type  = 'col'
    chart.title = f'Trámites por unidad — {periodo}'
    chart.style = 10
    data   = Reference(ws, min_col=2, max_col=4, min_row=1, max_row=ws.max_row)
    cats   = Reference(ws, min_col=1, min_row=2, max_row=ws.max_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.width  = 22
    chart.height = 14
    ws.add_chart(chart, 'G2')

    return _excel_response(wb, f'kpi_unidades_{timezone.now().strftime("%Y%m%d")}.xlsx')